from datetime import datetime, timedelta
from decimal import Decimal

from django.db.models import Sum, Count, Q, F
from django.http import HttpResponse
from rest_framework.viewsets import ViewSet
from rest_framework.decorators import action
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill

from common.permissions import DataIsolationMixin
from common.utils import success_response, error_response, get_client_ip

from apps.budget.models import Budget, BudgetStatus, BudgetItem, BudgetSource
from apps.department.models import Department, DepartmentType
from apps.workflow.models import ApprovalFlow, ApprovalStatus
from apps.audit.models import AuditLog


class ReportViewSet(ViewSet):
    """报表分析 / Dashboard 数据接口"""

    permission_classes = [IsAuthenticated]

    def _get_budget_qs(self, request, year=None):
        """根据权限获取预算查询集"""
        user = request.user
        queryset = Budget.objects.all()
        if year:
            queryset = queryset.filter(year=year)

        if user.is_first_budget_admin or user.is_first_budget_host or user.is_first_dept_head:
            pass
        elif user.is_second_budget_admin_primary or user.is_second_budget_admin_secondary or user.is_second_dept_head:
            if user.department:
                queryset = queryset.filter(department=user.department)
        else:
            queryset = queryset.none()
        return queryset

    @action(detail=False, methods=['get'])
    def dashboard(self, request):
        """
        GET /api/analysis/dashboard/ - 仪表盘总览数据（智慧城市风格）
        """
        year = int(request.query_params.get('year', datetime.now().year))
        budget_qs = self._get_budget_qs(request, year)

        # 按来源统计
        source_stats = budget_qs.values('source').annotate(
            total=Sum('total_amount'),
            count=Count('id')
        )
        source_data = {}
        for s in source_stats:
            source_data[s['source']] = {
                'total': str(s['total'] or 0),
                'count': s['count']
            }

        # 按部门统计（二级部门）
        dept_stats = budget_qs.filter(
            department__dept_type=DepartmentType.SECOND
        ).values('department__name').annotate(
            total=Sum('total_amount'),
            count=Count('id')
        ).order_by('-total')[:10]

        # 按类别统计 OPEX/CAPEX
        category_stats = budget_qs.values('category').annotate(
            total=Sum('total_amount'),
            count=Count('id')
        )
        category_data = {}
        for c in category_stats:
            category_data[c['category']] = {
                'total': str(c['total'] or 0),
                'count': c['count']
            }

        # 待审批统计
        pending_count = budget_qs.filter(status=BudgetStatus.PENDING).count()

        # 版本统计
        version_count = budget_qs.exclude(parent_version=None).count()

        # 总预算金额
        total_budget = budget_qs.aggregate(total=Sum('total_amount'))['total'] or 0

        return Response({
            'year': year,
            'overview': {
                'total_budget': str(total_budget),
                'budget_count': budget_qs.count(),
                'pending_count': pending_count,
                'version_count': version_count,
            },
            'source_distribution': source_data,
            'category_distribution': category_data,
            'department_ranking': [
                {
                    'department_name': d['department__name'],
                    'total': str(d['total'] or 0),
                    'count': d['count']
                } for d in dept_stats
            ],
        })

    @action(detail=False, methods=['get'])
    def overview(self, request):
        """
        GET /api/analysis/overview/ - 智慧城市风格总览
        """
        year = int(request.query_params.get('year', datetime.now().year))
        budget_qs = self._get_budget_qs(request, year)

        # 统计各状态数量
        status_counts = budget_qs.values('status').annotate(count=Count('id'))
        status_data = {s['status']: s['count'] for s in status_counts}

        # 按月度统计新增预算
        monthly_data = []
        for month in range(1, 13):
            count = budget_qs.filter(
                created_at__year=year,
                created_at__month=month
            ).count()
            monthly_data.append({'month': month, 'count': count})

        # SS Public 占比
        ss_public_total = budget_qs.filter(source=BudgetSource.SS_PUBLIC).aggregate(
            total=Sum('total_amount')
        )['total'] or 0

        return Response({
            'year': year,
            'status_distribution': status_data,
            'monthly_trend': monthly_data,
            'ss_public_ratio': {
                'amount': str(ss_public_total),
                'total': str(budget_qs.aggregate(total=Sum('total_amount'))['total'] or 0),
            }
        })

    @action(detail=False, methods=['get'], url_path='department-ranking')
    def department_ranking(self, request):
        """
        GET /api/analysis/department-ranking/ - 部门排名
        """
        year = int(request.query_params.get('year', datetime.now().year))
        limit = int(request.query_params.get('limit', 10))
        budget_qs = self._get_budget_qs(request, year)

        dept_stats = budget_qs.filter(
            department__dept_type=DepartmentType.SECOND
        ).values('department__id', 'department__name').annotate(
            total_amount=Sum('total_amount'),
            budget_count=Count('id')
        ).order_by('-total_amount')[:limit]

        result = []
        for d in dept_stats:
            result.append({
                'department_id': str(d['department__id']),
                'department_name': d['department__name'],
                'total_amount': str(d['total_amount'] or 0),
                'budget_count': d['budget_count'],
            })

        return Response(result)

    @action(detail=False, methods=['get'], url_path='monthly-trend')
    def monthly_trend(self, request):
        """
        GET /api/analysis/monthly-trend/ - 月度趋势
        """
        year = int(request.query_params.get('year', datetime.now().year))
        budget_qs = self._get_budget_qs(request, year)

        monthly_data = []
        for month in range(1, 13):
            month_qs = budget_qs.filter(created_at__year=year, created_at__month=month)
            stats = month_qs.aggregate(
                count=Count('id'),
                total=Sum('total_amount')
            )
            monthly_data.append({
                'month': month,
                'month_name': f'{month}月',
                'count': stats['count'] or 0,
                'total': str(stats['total'] or 0),
            })

        return Response(monthly_data)

    @action(detail=False, methods=['get'], url_path='category-analysis')
    def category_analysis(self, request):
        """
        GET /api/analysis/category-analysis/ - 类别分析（OPEX/CAPEX）
        """
        year = int(request.query_params.get('year', datetime.now().year))
        budget_qs = self._get_budget_qs(request, year)

        category_stats = budget_qs.values('category').annotate(
            total=Sum('total_amount'),
            count=Count('id')
        ).order_by('-total')

        result = []
        for stat in category_stats:
            result.append({
                'category': stat['category'],
                'total': str(stat['total'] or 0),
                'count': stat['count'],
            })

        return Response(result)

    @action(detail=False, methods=['get'], url_path='budget-summary')
    def budget_summary(self, request):
        """
        GET /api/analysis/budget-summary/ - 预算汇总（按部门+来源）
        """
        year = int(request.query_params.get('year', datetime.now().year))
        budget_qs = self._get_budget_qs(request, year)

        # 按部门和来源汇总
        summary = budget_qs.values(
            'department__id', 'department__name', 'source'
        ).annotate(
            total=Sum('total_amount'),
            count=Count('id')
        ).order_by('department__name', 'source')

        # 重组数据结构
        dept_map = {}
        for item in summary:
            dept_id = str(item['department__id'])
            if dept_id not in dept_map:
                dept_map[dept_id] = {
                    'department_id': dept_id,
                    'department_name': item['department__name'],
                    'sources': {}
                }
            dept_map[dept_id]['sources'][item['source']] = {
                'total': str(item['total'] or 0),
                'count': item['count']
            }

        return Response(list(dept_map.values()))

    @action(detail=False, methods=['get'], url_path='budget-usage')
    def budget_usage(self, request):
        """
        GET /api/analysis/budget-usage/ - 预算版本追踪列表
        """
        year = int(request.query_params.get('year', datetime.now().year))
        budget_qs = self._get_budget_qs(request, year)

        # 按部门获取最新审批通过的版本
        from apps.department.models import Department
        departments = Department.objects.filter(dept_type=DepartmentType.SECOND)

        result = []
        for dept in departments:
            latest = budget_qs.filter(
                department=dept,
                status=BudgetStatus.APPROVED
            ).order_by('-version_major', '-version_minor').first()

            if latest:
                result.append({
                    'budget_id': str(latest.id),
                    'budget_no': latest.budget_no,
                    'department_name': dept.name,
                    'year': latest.year,
                    'version': latest.version_label,
                    'total_amount': str(latest.total_amount),
                    'item_count': latest.items.filter(is_deleted=False).count(),
                    'source': latest.source,
                    'category': latest.category,
                })

        return Response(result)

    @action(detail=False, methods=['get'], url_path='export/budgets')
    def export_budgets(self, request):
        """
        GET /api/analysis/export/budgets/ - 导出预算报表 Excel
        """
        year = int(request.query_params.get('year', datetime.now().year))
        budget_qs = self._get_budget_qs(request, year)

        wb = Workbook()
        ws = wb.active
        ws.title = "预算报表"

        headers = ['预算编号', '部门', '类型', '来源', '年度', '版本', '总额', '条目数', '状态', '编制时间']
        ws.append(headers)

        header_fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
        header_font = Font(bold=True, color='FFFFFF')
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')

        for budget in budget_qs:
            ws.append([
                budget.budget_no or '-',
                budget.department.name if budget.department else '',
                budget.category,
                budget.get_source_display(),
                budget.year,
                budget.version_label,
                float(budget.total_amount or 0),
                budget.items.filter(is_deleted=False).count(),
                budget.get_status_display(),
                budget.created_at.strftime('%Y-%m-%d %H:%M') if budget.created_at else '',
            ])

        for col in ['A', 'B', 'C', 'D', 'E', 'F', 'G', 'H', 'I', 'J']:
            ws.column_dimensions[col].width = 15
        ws.column_dimensions['B'].width = 20

        user = request.user
        AuditLog.objects.create(
            user_id=str(user.id),
            user_name=user.name or user.username,
            action='EXPORT',
            module='report',
            target_type='budget_report',
            ip=get_client_ip(request),
        )

        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="budget_report.xlsx"'
        wb.save(response)
        return response

    @action(detail=False, methods=['get'], url_path='export/analysis')
    def export_analysis(self, request):
        """
        GET /api/analysis/export/analysis/ - 导出分析报表 Excel
        """
        user = request.user
        year = int(request.query_params.get('year', datetime.now().year))
        budget_qs = self._get_budget_qs(request, year)

        wb = Workbook()

        # 部门排名表
        ws1 = wb.active
        ws1.title = "部门排名"

        dept_stats = budget_qs.filter(
            department__dept_type=DepartmentType.SECOND
        ).values('department__name').annotate(
            total=Sum('total_amount'),
            count=Count('id')
        ).order_by('-total')[:20]

        ws1.append(['部门', '预算总额', '预算数'])
        for d in dept_stats:
            ws1.append([
                d['department__name'],
                float(d['total'] or 0),
                d['count'],
            ])

        # 类别分析表
        ws2 = wb.create_sheet("类别分析")
        category_stats = budget_qs.values('category').annotate(
            total=Sum('total_amount'),
            count=Count('id')
        ).order_by('-total')

        ws2.append(['类别', '总额', '预算数'])
        for c in category_stats:
            ws2.append([
                c['category'],
                float(c['total'] or 0),
                c['count'],
            ])

        # 来源分析表
        ws3 = wb.create_sheet("来源分析")
        source_stats = budget_qs.values('source').annotate(
            total=Sum('total_amount'),
            count=Count('id')
        ).order_by('-total')

        ws3.append(['来源', '总额', '预算数'])
        for s in source_stats:
            ws3.append([
                s['source'],
                float(s['total'] or 0),
                s['count'],
            ])

        for ws in [ws1, ws2, ws3]:
            header_fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
            header_font = Font(bold=True, color='FFFFFF')
            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font

        AuditLog.objects.create(
            user_id=str(user.id),
            user_name=user.name or user.username,
            action='EXPORT',
            module='report',
            target_type='analysis_report',
            ip=get_client_ip(request),
        )

        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="analysis_report.xlsx"'
        wb.save(response)
        return response
