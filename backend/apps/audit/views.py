from datetime import datetime, timedelta

from django.db.models import Count
from django.http import HttpResponse
from rest_framework import status
from rest_framework.viewsets import ReadOnlyModelViewSet
from rest_framework.decorators import action
from rest_framework.permissions import IsAuthenticated
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill

from common.pagination import StandardPagination
from common.permissions import IsAdmin
from common.utils import success_response, error_response, get_client_ip

from .models import AuditLog
from .serializers import AuditLogSerializer


class AuditLogViewSet(ReadOnlyModelViewSet):
    """审计日志查询（仅管理员）"""

    serializer_class = AuditLogSerializer
    permission_classes = [IsAuthenticated, IsAdmin]

    def get_queryset(self):
        qs = AuditLog.objects.all()
        params = self.request.query_params

        # 用户筛选
        if params.get('userId'):
            qs = qs.filter(user_id=params['userId'])

        # 模块筛选
        if params.get('module'):
            qs = qs.filter(module=params['module'])

        # 操作类型筛选
        if params.get('action'):
            qs = qs.filter(action=params['action'])

        # 时间范围筛选
        if params.get('startDate'):
            qs = qs.filter(created_at__gte=params['startDate'])
        if params.get('endDate'):
            qs = qs.filter(created_at__lte=params['endDate'])

        # 目标类型筛选
        if params.get('targetType'):
            qs = qs.filter(target_type=params['targetType'])

        # 目标ID筛选
        if params.get('targetId'):
            qs = qs.filter(target_id=params['targetId'])

        return qs

    @action(detail=False, methods=['get'])
    def stats(self, request):
        """
        GET /api/audit-logs/stats/ - 审计日志统计信息

        返回:
            - by_module: 按模块统计数量
            - by_action: 按操作类型统计
            - daily_trend: 最近7天趋势
        """
        # 按模块统计
        by_module = list(
            AuditLog.objects.values('module')
            .annotate(count=Count('id'))
            .order_by('-count')
        )

        # 按操作类型统计
        by_action = list(
            AuditLog.objects.values('action')
            .annotate(count=Count('id'))
            .order_by('-count')
        )

        # 最近7天趋势
        end_date = datetime.now()
        start_date = end_date - timedelta(days=6)

        daily_trend = []
        for i in range(7):
            date = (start_date + timedelta(days=i)).date()
            count = AuditLog.objects.filter(
                created_at__date=date
            ).count()
            daily_trend.append({
                'date': date.isoformat(),
                'count': count
            })

        return success_response({
            'by_module': by_module,
            'by_action': by_action,
            'daily_trend': daily_trend,
        })

    @action(detail=False, methods=['get'])
    def export(self, request):
        """
        GET /api/audit-logs/export/ - 导出审计日志为 Excel

        查询参数:
            - userId: 用户ID
            - module: 模块
            - action: 操作类型
            - startDate: 开始日期
            - endDate: 结束日期
            - targetType: 目标类型
            - targetId: 目标ID
        """
        queryset = self.get_queryset()

        # 创建 Excel 工作簿
        wb = Workbook()
        ws = wb.active
        ws.title = "审计日志"

        # 设置表头
        headers = ['ID', '用户名', '操作类型', '模块', '目标类型', '目标ID', 'IP地址', '操作时间']
        ws.append(headers)

        # 设置表头样式
        header_fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
        header_font = Font(bold=True, color='FFFFFF')
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')

        # 填充数据
        action_map = {
            'CREATE': '创建',
            'UPDATE': '更新',
            'DELETE': '删除',
            'APPROVE': '审批通过',
            'REJECT': '审批驳回',
            'SUBMIT': '提交',
            'ADJUST': '调整',
            'LOGIN': '登录',
            'LOGOUT': '登出',
            'EXPORT': '导出',
            'IMPORT': '导入',
        }
        module_map = {
            'budget': '预算管理',
            'purchase': '采购管理',
            'approval': '审批管理',
            'user': '用户管理',
            'department': '部门管理',
            'system': '系统管理',
            'report': '报表分析',
            'import_export': '导入导出',
        }

        for log in queryset:
            ws.append([
                str(log.id),
                log.user_name,
                action_map.get(log.action, log.action),
                module_map.get(log.module, log.module),
                log.target_type or '',
                log.target_id or '',
                log.ip or '',
                log.created_at.strftime('%Y-%m-%d %H:%M:%S'),
            ])

        # 调整列宽
        ws.column_dimensions['A'].width = 40
        ws.column_dimensions['B'].width = 15
        ws.column_dimensions['C'].width = 12
        ws.column_dimensions['D'].width = 12
        ws.column_dimensions['E'].width = 15
        ws.column_dimensions['F'].width = 40
        ws.column_dimensions['G'].width = 15
        ws.column_dimensions['H'].width = 20

        # 记录导出审计日志
        AuditLog.objects.create(
            user_id=str(request.user.id),
            user_name=request.user.name or request.user.username,
            action='EXPORT',
            module='system',
            target_type='audit_logs',
            ip=get_client_ip(request),
        )

        # 返回响应
        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="audit_logs.xlsx"'
        wb.save(response)
        return response
